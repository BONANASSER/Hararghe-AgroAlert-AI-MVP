"""
generate_six_month_plan_doc.py
Compiles the comprehensive 6-month master execution plan for Bona Nasser Abdo's 
MSc Thesis under the supervision of Dr. Million Meshesha at Haramaya University.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

HEX_PRIMARY = "1A365D"      # Deep Navy
HEX_SECONDARY = "2B6CB0"    # Slate Blue
HEX_ACCENT = "2C7A7B"       # Teal
HEX_DARK = "2D3748"         # Charcoal Text
HEX_LIGHT_BG = "F7FAFC"     # Off-white / light slate
HEX_CALLOUT_BG = "EDF2F7"   # Table / Callout Shading
HEX_BORDER = "CBD5E0"       # Border Grey
HEX_HIGHLIGHT = "C53030"    # Crimson for alerts

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, border_hex="CBD5E0"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(26, 54, 93)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(43, 108, 176)
    return p

def add_body(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(26, 54, 93)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(26, 54, 93)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_callout(doc, text, title="STRATEGIC DIRECTIVE / ADVISORY MILESTONE", alert_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_hex = "EDF2F7" if alert_type == "info" else "FFF5F5"
    border_hex = "2B6CB0" if alert_type == "info" else "E53E3E"
    
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}]\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(9)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(43, 108, 176) if alert_type == "info" else RGBColor(197, 48, 48)
    
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(45, 55, 72)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def style_table(tbl, col_widths, headers, data):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, "CBD5E0")
    
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        bg_color = HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = col_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(45, 55, 72)

def generate_plan_document():
    doc = docx.Document()
    
    # 0.8-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("MSc Thesis 6-Month Strategic Roadmap | Bona Nasser Abdo | Advisor: Dr. Million Meshesha")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(113, 128, 150)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Haramaya University | Postgraduate Directorate | Page ")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(113, 128, 150)
        add_page_number(frun)

    # Document Header
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    r1 = p_inst.add_run("HARAMAYA UNIVERSITY | COLLEGE OF COMPUTING AND INFORMATICS\nDEPARTMENT OF INFORMATION SCIENCE | POSTGRADUATE RESEARCH PROGRAM")
    r1.font.name = "Arial"
    r1.font.size = Pt(10)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(26, 54, 93)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(6)
    r_t = p_title.add_run("SIX-MONTH MASTER EXECUTION ROADMAP: FROM PROPOSAL TO FINAL DEFENSE\nAI-Based Multimodal Deep Learning Early Warning System for Drought, Crop Advisory, and Food Insecurity Prediction in Eastern Ethiopia")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(12)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(43, 108, 176)

    # Meta Table
    tbl_meta = doc.add_table(rows=2, cols=2)
    meta_headers = ["Research Metadata", "Operational Specification"]
    meta_data = [
        [("Graduate Researcher", "Bona Nasser Abdo (MSc Data Science Candidate)"), ("Principal Advisor", "Dr. Million Meshesha (Associate Professor)")],
        [("Total Time Horizon", "6 Calendar Months (24 Academic Weeks)"), ("Milestone Gates", "Gate 1: Proposal Defense (Month 2) | Gate 2: Final Thesis Defense (Month 6)")]
    ]
    for r_i, row in enumerate(meta_data):
        for c_i, (k, v) in enumerate(row):
            cell = tbl_meta.cell(r_i, c_i)
            set_cell_background(cell, "F7FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            rk = p.add_run(f"{k}: ")
            rk.font.name = "Calibri"
            rk.font.size = Pt(8.5)
            rk.font.bold = True
            rk.font.color.rgb = RGBColor(26, 54, 93)
            rv = p.add_run(v)
            rv.font.name = "Calibri"
            rv.font.size = Pt(8.5)
            rv.font.color.rgb = RGBColor(45, 55, 72)
    set_table_borders(tbl_meta, "CBD5E0")

    add_body(doc, "This strategic roadmap outlines the rigorous, week-by-week operational execution plan to complete, defend, and submit the MSc Data Science thesis within a strict six-month timeline. The lifecycle is divided into two major operational phases: Phase I (Proposal Development, Defense & Approval, Months 1–2) and Phase II (Empirical Implementation, Modeling, Writing & Final Thesis Defense, Months 3–6).", space_after=6)

    # MASTER GANTT TABLE
    add_heading_1(doc, "1. Master 24-Week Chronological Gantt Schedule")
    gantt_headers = ["Phase / Month", "Week Span", "Core Focus & Milestones", "Primary Deliverable / Milestone Gate"]
    gantt_widths = [Inches(1.2), Inches(0.9), Inches(2.7), Inches(1.7)]
    gantt_data = [
        ["Phase I: Month 1", "Weeks 1–4", "Literature review, 2-page problem brief for Dr. Million, theoretical framing, proposal drafting (Ch 1–3).", "Complete Proposal Draft (Chapters 1, 2, 3)"],
        ["Phase I: Month 2", "Weeks 5–8", "Advisor review, DGC submission, proposal slide deck preparation, public defense, post-defense revisions.", "GATE 1: Approved Thesis Proposal"],
        ["Phase II: Month 3", "Weeks 9–12", "Automated ingestion (NOAA, CHIRPS, SMAP, MODIS, ERA5), WGS84 woreda alignment, Sc-PDSI climatology engine.", "Harmonized Woreda × Month Data Cube"],
        ["Phase II: Month 4", "Weeks 13–16", "Conv1D-LSTM ENSO model, ResNet-50 patch CNN, LSTM sequence, FFT harmonic, Tabular MLP, Gated Cross-Attention.", "Trained Model Checkpoints & Baselines"],
        ["Phase II: Month 5", "Weeks 17–20", "Explainable AI (SHAP, Grad-CAM), Neo4j knowledge graph, 4-tier advisory matrix, out-of-sample testing (2024–26).", "Empirical Results, XAI & Decision App"],
        ["Phase II: Month 6", "Weeks 21–24", "Full manuscript compilation (Ch 1–5), pre-defense review with Dr. Million, final public defense, hardbound submission.", "GATE 2: Final Defended MSc Thesis"]
    ]
    tbl_gantt = doc.add_table(rows=1, cols=4)
    style_table(tbl_gantt, gantt_widths, gantt_headers, gantt_data)

    # DETAILED BREAKDOWN
    add_heading_1(doc, "2. Detailed Step-by-Step Weekly Execution Plan")

    # PHASE I
    add_heading_2(doc, "PHASE I: PROPOSAL DEVELOPMENT, DEFENSE & APPROVAL (MONTHS 1–2 / WEEKS 1–8)")
    
    add_heading_2(doc, "Month 1: Problem Formulation, Literature Syntheses & Proposal Drafting")
    add_bullet(doc, "Synthesize state-of-the-art literature (Busker et al., 2024; Baye et al., 2026; Abate et al., 2025; Megbar & Tadesse, 2016). Formulate theoretical framework bridging physical climatology to acute food insecurity.", bold_prefix="Week 1 (Literature & Theoretical Grounding): ")
    add_bullet(doc, "Finalize and submit the formal 2-page Statement of the Problem and Data Availability Verification to Dr. Million Meshesha. Incorporate advisor feedback into Chapter 1 (Background, Problem, 3 Research Questions, Objectives).", bold_prefix="Week 2 (Statement of Problem & Advisor Alignment): ")
    add_bullet(doc, "Draft Chapter 2 (Literature Review & Conceptual Framework) and draft Chapter 3 (Study Area profile of Eastern Hararghe, 5-tier architecture, mathematical Sc-PDSI formulations, and Super-Hybrid model design).", bold_prefix="Week 3 (Methodology & Architecture Drafting): ")
    add_bullet(doc, "Verify zero-credential data ingestion scripts (`data/raw/manifests/dataset.json` with SHA-256 hashes). Assemble complete proposal manuscript (Chapters 1–3, References, Work Plan, Budget).", bold_prefix="Week 4 (Proposal Assembly & Data Audit): ")

    add_heading_2(doc, "Month 2: Advisor Iteration, Departmental Submission & Proposal Defense")
    add_bullet(doc, "Formal review session with Dr. Million Meshesha. Refine methodology based on supervisory critique (e.g., target label definitions, class imbalance handling, validation boundaries).", bold_prefix="Week 5 (Advisory Review & One-on-One Critiques): ")
    add_bullet(doc, "Incorporate all corrections. Submit formal proposal copies and approval sheets to the Department Graduate Committee (DGC) of Information Science.", bold_prefix="Week 6 (Formal DGC Submission & Scheduling): ")
    add_bullet(doc, "Develop professional defense slide deck (20 slides: Problem, Research Gaps, 3 RQs, 5-Tier Architecture, Sc-PDSI Physics, Validation). Defend thesis proposal before departmental academic panel.", bold_prefix="Week 7 (Proposal Defense Presentation - GATE 1): ")
    add_bullet(doc, "Incorporate examiner comments. Secure official approval signatures from advisors and DGC Chairperson. Obtain formal institutional clearance.", bold_prefix="Week 8 (Post-Defense Revisions & Institutional Clearance): ")

    add_callout(doc, "MILESTONE GATE 1 ACHIEVED: Officially approved and defended MSc Thesis Proposal registered with the Postgraduate Program Directorate of Haramaya University. Proceed to empirical execution.", title="GATE 1 MILESTONE: PROPOSAL DEFENSE & APPROVAL", alert_type="info")

    # PHASE II
    add_heading_2(doc, "PHASE II: EMPIRICAL IMPLEMENTATION, EXPERIMENTAL RIGOR & FINAL DEFENSE (MONTHS 3–6)")

    add_heading_2(doc, "Month 3: Automated Ingestion, Harmonization & Physical Climatology Engine")
    add_bullet(doc, "Deploy automated data connectors downloading verified archives: NOAA PSL (Niño 3.4), CHIRPS v3 (precip), NASA SMAP (soil moisture), MODIS MOD13Q1/MOD11A2, ERA5, SoilGrids, and FEWS NET IPC shapefiles. Verify SHA-256 manifests.", bold_prefix="Week 9 (Automated Ingestion & Manifest Security): ")
    add_bullet(doc, "Reproject all spatial grids to WGS84 (EPSG:4326). Perform zonal statistics across Eastern Hararghe woreda polygons. Slice standardized 32×32 pixel multi-spectral visual patches centered over woreda centroids.", bold_prefix="Week 10 (Spatial Alignment & Patch Extraction): ")
    add_bullet(doc, "Code and calibrate the physical two-layer Self-Calibrated Palmer Drought Severity Index (Sc-PDSI) with CAFEC precipitation and Z-index moisture departures. Compute SPI-1, SPI-3, SPI-6, and VCI rasters.", bold_prefix="Week 11 (Physical Climatology & Water Balance Modeling): ")
    add_bullet(doc, "Assemble master analytical table structured as Woreda × Month. Generate 1-, 2-, 3-, and 6-month historical lags strictly using antecedent observations prior to timestamp t (zero future data leakage).", bold_prefix="Week 12 (Master Analytical Data Cube Assembly): ")

    add_heading_2(doc, "Month 4: Deep Learning Development, Baseline Training & Fusion Modeling")
    add_bullet(doc, "Train and calibrate the Conv1D-LSTM oceanic teleconnection forecaster. Evaluate Niño 3.4 SST anomaly predictions across 1–12 leads against persistence and ridge baselines.", bold_prefix="Week 13 (Conv1D-LSTM ENSO Teleconnection Modeling): ")
    add_bullet(doc, "Build and pre-train the four specialized encoders: (a) ResNet-50 spatial patch CNN, (b) 12-month recurrent LSTM temporal sequence branch, (c) Fast Fourier Transform (FFT) harmonic spectral branch, and (d) structured tabular MLP.", bold_prefix="Week 14 (Specialized Modality Encoder Construction): ")
    add_bullet(doc, "Implement the Multimodal Gated Cross-Attention Fusion Engine. Connect multi-task heads for Drought Hazard P(SPI-3<=-1), Crop Risk P_adj, and FEWS NET / IPC Acute Food Insecurity (Phases 1–5). Apply Focal Loss (gamma=2.0) and class-weighted cross-entropy.", bold_prefix="Week 15 (Cross-Attention Fusion & Multi-Task Training): ")
    add_bullet(doc, "Train classical machine learning baselines (Logistic Regression, Random Forest, XGBoost) and unimodal deep networks under identical chronological splits (2015–2021 train, 2022–2023 validation). Record comparative performance tables.", bold_prefix="Week 16 (Baseline Training & Comparative Benchmarking): ")

    add_heading_2(doc, "Month 5: Explainable AI, Knowledge Graph, Decision Platform & Evaluation")
    add_bullet(doc, "Implement Explainable AI pipelines: compute SHAP feature attributions (global bar charts & local waterfall plots) for tabular features; generate Grad-CAM heatmaps for satellite visual patches; compute temporal lag attributions.", bold_prefix="Week 17 (Explainable AI - SHAP & Grad-CAM Implementation): ")
    add_bullet(doc, "Construct the Neo4j Semantic Knowledge Graph. Ingest woredas, crops, climate observations, hazard events, and IPC phases. Implement Cypher queries for multi-hazard cascading vulnerability reasoning.", bold_prefix="Week 18 (Neo4j Semantic Knowledge Graph Modeling): ")
    add_bullet(doc, "Calibrate biological crop sensitivity multipliers (s) for sorghum, maize, haricot bean, wheat, and tef. Map adjusted risks to the 4-tier prescriptive advisory matrix. Implement the Bi-Directional IEK Consensus Engine.", bold_prefix="Week 19 (Prescriptive Advisory & IEK Consensus Engine): ")
    add_bullet(doc, "Execute independent out-of-sample hindcasting on the strictly withheld 2024–2026 test window. Compute Brier Skill Scores (BSS), ROC-AUC, PR-AUC, Macro F1, and reliability calibration curves. Enforce the 'Honest Lead Horizon' principle.", bold_prefix="Week 20 (Independent Hindcasting & Honest Lead Validation): ")

    add_heading_2(doc, "Month 6: Manuscript Compilation, Advisory Review & Final Thesis Defense")
    add_bullet(doc, "Compile complete Master Thesis manuscript: Chapter 1 (Introduction), Chapter 2 (Literature Review), Chapter 3 (Methodology), Chapter 4 (Results & In-Depth Discussion with 17+ tables and figures), Chapter 5 (Conclusions & Recommendations), and Appendices A–E.", bold_prefix="Week 21 (Complete Thesis Manuscript Drafting): ")
    add_bullet(doc, "Comprehensive manuscript review with Dr. Million Meshesha. Conduct mock defense presentations. Revise all chapters based on detailed advisory critique.", bold_prefix="Week 22 (Advisory Review, Revisions & Pre-Defense Clearance): ")
    add_bullet(doc, "Deliver formal public defense presentation before Internal Examiner, External Examiner, DGC Chairperson, and academic community. Present live decision-support dashboard and empirical validation results.", bold_prefix="Week 23 (FINAL MSC THESIS DEFENSE - GATE 2): ")
    add_bullet(doc, "Incorporate minor examiner corrections. Obtain official defense approval signatures. Prepare final hardbound copies and electronic submission to Haramaya University Library and Institutional Repository.", bold_prefix="Week 24 (Final Corrections, Sign-Off & Institutional Submission): ")

    add_callout(doc, "MILESTONE GATE 2 ACHIEVED: Final Master of Science degree requirements fulfilled. Completed, defended, and submitted MSc Data Science Thesis with open-source codebase and peer-reviewed journal draft ready for submission.", title="GATE 2 MILESTONE: DEGREE COMPLETION & FINAL SUBMISSION", alert_type="info")

    # RISK MANAGEMENT TABLE
    add_heading_1(doc, "3. Technical & Academic Risk Mitigation Matrix")
    risk_headers = ["Potential Execution Risk", "Impact / Probability", "Proactive Mitigation Strategy"]
    risk_widths = [Inches(1.8), Inches(1.3), Inches(3.4)]
    risk_data = [
        ["GPU compute bottlenecks or out-of-memory errors during multimodal patch training", "High / Medium", "Utilize mixed-precision (FP16) training, gradient accumulation, and batch size auto-scaling; leverage Google Colab Pro / Kaggle GPU instances."],
        ["Delayed or missing food-insecurity labels from official agencies", "High / Low", "Supervised food-security modeling is strictly bounded to verified 2020–2026 FEWS NET shapefiles; physical Sc-PDSI/SPI targets cover the full 2015–2026 window."],
        ["Advisor availability constraints or delayed review cycles", "Medium / Medium", "Maintain strict dekadal progress memos, shared Overleaf/Word cloud documents, and bi-weekly milestone check-ins with Dr. Million Meshesha."],
        ["Class imbalance distorting high-risk drought and IPC Phase 4/5 predictions", "High / High", "Apply Focal Loss (gamma=2.0), class-weighted cross-entropy, and SMOTE exclusively within training folds; evaluate via PR-AUC and Macro F1."]
    ]
    tbl_risk = doc.add_table(rows=1, cols=3)
    style_table(tbl_risk, risk_widths, risk_headers, risk_data)

    output_path = os.path.join(os.getcwd(), "MSc_Thesis_6_Month_Master_Execution_Plan.docx")
    doc.save(output_path)
    print(f"6-Month Master Plan successfully created at: {output_path}")

if __name__ == "__main__":
    generate_plan_document()
