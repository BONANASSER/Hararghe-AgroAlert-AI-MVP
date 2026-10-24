"""
build_chapter_2.py
Chapter 2: Theoretical Foundations and Conceptual Framework
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_master_concept_note import (
    add_heading_1, add_heading_2, add_heading_3, add_body, add_bullet, 
    add_callout, style_table, HEX_PRIMARY, HEX_SECONDARY, HEX_CALLOUT_BG
)

def build_chapter_2(doc):
    add_heading_1(doc, "CHAPTER TWO: THEORETICAL FOUNDATIONS AND CONCEPTUAL FRAMEWORK")
    
    # 2.1 Causal Continuum
    add_heading_2(doc, "2.1 The Hydro-Climatic to Food Insecurity Causal Continuum")
    add_body(doc, "The genesis of drought-induced acute food insecurity is neither an instantaneous nor a purely meteorological event. It represents a complex, non-linear cascading process propagating through distinct physical, ecological, agronomical, and socioeconomic subsystems. Developing an effective multimodal early warning system requires grounding artificial intelligence within this physical causal continuum:")
    
    add_bullet(doc, "Large-scale sea surface temperature (SST) anomalies in the equatorial Pacific Ocean (El Niño–Southern Oscillation [ENSO]) and equatorial Indian Ocean (Indian Ocean Dipole [IOD]) alter the global atmospheric Walker and Hadley circulations. Anomalous eastward displacement of deep tropical convection in the Pacific triggers large-scale atmospheric subsidence and high-pressure anomalies over the Ethiopian Highlands, suppressing local convective precipitation.", bold_prefix="1. Macro-Scale Oceanic-Atmospheric Forcing: ")
    add_bullet(doc, "Suppressed atmospheric moisture advection results in severe delays or failures of the bimodal rainfall regimes (Belg and Kiremt / Deyr) in Eastern Ethiopia, generating high vapor pressure deficits and extreme potential evapotranspiration (PET).", bold_prefix="2. Meteorological Moisture Deficit: ")
    add_bullet(doc, "Persistent rainfall deficits trigger progressive drawdown of the soil moisture profile. Depletion initially impacts the immediate surface layer (0–10 cm) before exhausting available water capacity (AWC) in the deeper root-zone (10–100 cm), terminating capillary recharge.", bold_prefix="3. Hydrological & Soil Moisture Depletion: ")
    add_bullet(doc, "Soil moisture exhaustion forces plants into permanent wilting point conditions. Transpiration ceases, chlorophyll degradation accelerates, and satellite sensors register sharp declines in Normalized Difference Vegetation Index (NDVI) and Vegetation Condition Index (VCI), accompanied by elevated Land Surface Temperatures (LST).", bold_prefix="4. Agricultural Crop Desiccation: ")
    add_bullet(doc, "Moisture stress during critical crop phenological stages (crown root initiation in cereals, tasseling/silking in maize, flowering and podding in haricot beans) causes catastrophic yield collapse or total crop abortion. Concurrently, grazing pasture degenerates, leading to livestock emaciation and mortality.", bold_prefix="5. Agronomic Production Failure: ")
    add_bullet(doc, "Production failure triggers severe supply shortages in local rural markets, causing staple food price hyper-inflation. Smallholder households experience catastrophic collapses in agricultural income, deterioration of terms-of-trade (livestock-to-grain price ratios), and rapid depletion of household coping assets.", bold_prefix="6. Socioeconomic & Market Disruption: ")
    add_bullet(doc, "Households exhaust adaptive strategies, resulting in severe food consumption deficits, acute malnutrition (wasting), distress migration, and escalation into severe humanitarian phases (IPC Phase 3 Crisis, Phase 4 Emergency, or Phase 5 Catastrophe/Famine).", bold_prefix="7. Acute Human Food Insecurity: ")

    # 2.2 Physical Climatology & Formulations
    add_heading_2(doc, "2.2 Physical Climatology and Mathematical Formulations of Drought")
    add_body(doc, "Rather than treating drought as a purely statistical machine-learning abstraction, this research couples empirical deep learning with rigorous physical hydrology. We formulate and implement three foundational hydro-climatic modeling frameworks:")
    
    add_heading_3(doc, "2.2.1 Self-Calibrated Palmer Drought Severity Index (Sc-PDSI)")
    add_body(doc, "The Palmer Drought Severity Index (PDSI) models soil-water accounting across a two-layer soil profile. In this study, the surface layer has an available water capacity (Ss) of 25.4 mm (1.0 inch). The underlying sub-surface layer holds the remainder of total available water capacity (Su = AWC - Ss), defaulting to 100 mm based on regional SoilGrids edaphic surveys. The monthly water balance is governed by:")
    
    add_callout(doc, 
        "ΔS = R - L\n"
        "P = ET + RO + ΔS = ET + RO + R - L\n\n"
        "Where:\n"
        "• P is observed monthly precipitation (mm)\n"
        "• ET is actual evapotranspiration (mm)\n"
        "• RO is surface and subsurface runoff (mm)\n"
        "• R is soil moisture recharge to the two layers (mm)\n"
        "• L is total soil moisture loss/drawdown (mm)\n"
        "• ΔS is the net change in stored profile soil moisture (mm)\n\n"
        "CAFEC Precipitation (Climatically Appropriate For Existing Conditions):\n"
        "P_hat_{m,g} = α_{m,g} * PET_{m,g} + β_{m,g} * PR_{m,g} + γ_{m,g} * PRO_{m,g} - δ_{m,g} * PL_{m,g}\n\n"
        "Where potential evapotranspiration (PET), potential recharge (PR), potential runoff (PRO), and potential loss (PL) are scaled by calibration coefficients α, β, γ, δ fitted separately for calendar month m and grid cell g across the 1990–2026 calibration record:\n"
        "α_{m,g} = Σ(ET) / Σ(PET),   β_{m,g} = Σ(R) / Σ(PR),   γ_{m,g} = Σ(RO) / Σ(PRO),   δ_{m,g} = Σ(L) / Σ(PL)",
        title="MATHEMATICAL FORMULATION: SC-PDSI WATER BALANCE & CAFEC PRECIPITATION", alert_type="info"
    )
    
    add_body(doc, "The moisture anomaly d is computed as the departure between observed precipitation and CAFEC precipitation, weighted by Palmer's climatic characteristic K to yield the Z-index (moisture anomaly index):")
    
    add_callout(doc,
        "d_{m,g} = P_{m,g} - P_hat_{m,g}\n"
        "Z_{m,g} = K_{m,g} * d_{m,g}\n\n"
        "Where K is the climatic weighting factor derived from calibration-period mean departures d_bar and moisture flux ratio Q:\n"
        "K'_{m,g} = 1.5 * log10[(Q_{m,g} + 2.8) / d_bar_{m,g}] + 0.5\n"
        "K_{m,g} = K'_{m,g} * (17.67 / Σ(d_bar * K'))\n\n"
        "Palmer Severity Recursion and Dynamic Self-Calibration:\n"
        "X_{t,raw} = 0.897 * X_{t-1,raw} + (Z_{t} / 3)\n\n"
        "In the project self-calibration procedure, extreme severity tails are rescaled against historical 2nd percentile (q_{0.02}) and 98th percentile (q_{0.98}) calibration records at each cell g:\n"
        "Sc-PDSI_{t,g} = -4.0 * (X_{t,raw} / |q_{0.02,g}|)   [when X_{t,raw} < 0]\n"
        "Sc-PDSI_{t,g} = +4.0 * (X_{t,raw} / q_{0.98,g})     [when X_{t,raw} >= 0]",
        title="MATHEMATICAL FORMULATION: Z-INDEX, RECURSION & DYNAMIC SELF-CALIBRATION", alert_type="info"
    )

    add_heading_3(doc, "2.2.2 Standardized Precipitation Index (SPI)")
    add_body(doc, "The Standardized Precipitation Index (SPI) is computed across multiple accumulation timescales (1-month SPI-1 for immediate meteorological stress, 3-month SPI-3 for agricultural root-zone moisture, and 6-month SPI-6 for hydrological reservoir balance). Long-term monthly precipitation totals are fitted to a two-parameter Gamma probability density function g(x):")
    add_body(doc, "g(x) = [1 / (β^α * Γ(α))] * x^(α - 1) * e^(-x / β)   for x > 0\nWhere α > 0 is the shape parameter, β > 0 is the scale parameter, and Γ(α) is the Gamma function. Cumulative probability distributions are transformed to standard normal distributions Z ~ N(0, 1), yielding SPI values where SPI <= -1.0 denotes moderate drought, SPI <= -1.5 denotes severe drought, and SPI <= -2.0 denotes extreme drought.")

    add_heading_3(doc, "2.2.3 Vegetation Condition Index (VCI)")
    add_body(doc, "To separate weather-induced physiological vegetation stress from static ecological landscape variations, the satellite-derived Vegetation Condition Index (VCI) is formulated from 16-day MODIS MOD13Q1 NDVI rasters:")
    add_body(doc, "VCI_{i,g} = [(NDVI_{i,g} - NDVI_{min,g}) / (NDVI_{max,g} - NDVI_{min,g})] * 100%\nWhere NDVI_{i,g} is the observed composite greenness for period i, and NDVI_{min,g} and NDVI_{max,g} represent the absolute multi-year historical minimum and maximum NDVI for that calendar window at grid cell g. VCI values below 35% indicate severe agricultural drought stress.")

    # 2.3 Atmospheric Teleconnections
    add_heading_2(doc, "2.3 Atmospheric Teleconnections: Oceanic-Atmospheric Forcing in Eastern Ethiopia")
    add_body(doc, "Rainfall variability across Eastern Ethiopia is teleconnected to sea surface temperature (SST) anomalies in the tropical Pacific and Indian Oceans through atmospheric Walker Circulation perturbations. Under El Niño conditions (warm SST anomalies in the Niño 3.4 domain [5°N–5°S, 170°W–120°W]), anomalous eastward displacement of tropical convective cells induces strong atmospheric subsidence and suppressed convection over the Ethiopian Highlands. Conversely, a positive Indian Ocean Dipole (+IOD: warm western Indian Ocean, cool eastern Indian Ocean) frequently enhances moisture flux toward East Africa during the Deyr / autumn season, though severe negative IOD (-IOD) phases amplify devastating multi-season droughts.")
    add_body(doc, "In this research, Pacific and Indian Ocean teleconnections are mathematically quantified using cross-correlation lag analysis spanning 0 to 9 month offsets across four Pacific indices (Niño 3.4, Niño 1+2, Niño 4, Southern Oscillation Index [SOI]) and the Dipole Mode Index (DMI). A dedicated Conv1D-LSTM neural forecaster is designed to predict Niño 3.4 SST trajectories up to 12 months in advance, providing the deep learning system with antecedent global boundary forcing.")

    # 2.4 Deep Learning & Multimodal Paradigms
    add_heading_2(doc, "2.4 Deep Learning and Multimodal Representation Learning Paradigms")
    add_body(doc, "Agricultural drought and food insecurity prediction present a multi-faceted spatiotemporal challenge where data structures differ fundamentally. Unimodal models fail because they cannot jointly represent spatial textures, sequential trajectories, cyclic harmonics, and structured tabular factors. We establish a multi-branch representation strategy:")
    
    add_bullet(doc, "2D Convolutional Neural Networks (CNNs) and deep residual networks (ResNet-50) excel at extracting localized spatial correlations, spatial gradients, and morphological textures from gridded satellite imagery (MODIS NDVI, EVI, and Daytime/Nighttime LST). Convolutions capture spatial contiguity and edge effects across topographic basins.", bold_prefix="Spatial Feature Representation (2D-CNN): ")
    add_bullet(doc, "Long Short-Term Memory (LSTM) recurrent neural networks capture long-range temporal dependencies, memory drawdown, and non-linear lag relationships across sequential hydro-climatic variables (CHIRPS precipitation, SMAP soil moisture, ERA5 temperature, and wind vectors), mitigating vanishing gradient bottlenecks.", bold_prefix="Temporal Sequence Representation (LSTM): ")
    add_bullet(doc, "Fast Fourier Transforms (FFT) map temporal climate series into the frequency domain, decomposing signals into fundamental sinusoidal frequencies. This harmonic branch encodes planetary astronomical periodicity, bi-annual Belg/Kiremt transitions, and multi-annual solar/ENSO cycles, anchoring the model against chaotic phase drift at longer lead times.", bold_prefix="Harmonic & Periodic Representation (FFT): ")
    add_bullet(doc, "Multilayer Perceptrons (MLPs) with batch normalization and dropout process structured tabular agricultural statistics (crop area, historical production), demographic density (WorldPop), and static soil edaphic properties (SoilGrids pH, clay, SOC).", bold_prefix="Structured Tabular Representation (MLP): ")
    add_bullet(doc, "Rather than naive vector concatenation (which allows high-dimensional image tensors to dominate low-dimensional tabular features), we introduce a Gated Cross-Attention Multimodal Fusion Engine. Cross-attention computes dynamic attention weights between spatial visual tokens, temporal latent states, and tabular vulnerability vectors, filtering irrelevant noise and amplifying cross-modal synergies.", bold_prefix="Multimodal Fusion Paradigms: ")

    # 2.5 Explainable AI
    add_heading_2(doc, "2.5 Explainable Artificial Intelligence (XAI) in Environmental Modeling")
    add_body(doc, "In humanitarian disaster management and agricultural advisory systems, 'black-box' deep learning models encounter deep skepticism from agronomists, extension agents, and policy makers. To ensure transparency, scientific defensibility, and algorithmic accountability, this thesis integrates three complementary Explainable AI (XAI) methodologies:")
    add_bullet(doc, "Based on cooperative game theory, SHAP computes the Shapley value of each tabular and fused feature, quantifying its positive or negative marginal contribution to individual woreda-level risk predictions. SHAP generates global feature importance rankings and local waterfall/force plots for every alert issued.", bold_prefix="1. SHAP (SHapley Additive exPlanations): ")
    add_bullet(doc, "Grad-CAM uses the gradient of the predicted risk score flowing into the final convolutional layer of the spatial CNN branch to produce coarse 2D heatmaps highlighting the exact geographical patches (e.g., specific valley corridors or deforested slopes) that triggered high-risk alerts.", bold_prefix="2. Grad-CAM (Gradient-Weighted Class Activation Mapping): ")
    add_bullet(doc, "Temporal attention weights and input perturbation analyses quantify which specific antecedent months (e.g., rainfall failure 2 months prior vs. soil moisture deficit 1 month prior) exerted the dominant influence on the forecast.", bold_prefix="3. Temporal Lag Attribution: ")

    # 2.6 Knowledge Graphs
    add_heading_2(doc, "2.6 Semantic Knowledge Graphs in Early Warning and Disaster Risk Reduction")
    add_body(doc, "While neural networks excel at continuous pattern recognition, they cannot explicitly represent relational hierarchies, institutional networks, or multi-hazard cascading impacts. We implement a semantic Knowledge Graph in Neo4j governed by a domain-specific agricultural ontology. The graph explicitly connects spatial entities (Woredas, Kebeles) to dynamic observations (ClimateRecords, SoilMoistureRecords, VegetationRecords), agricultural assets (Crops, Livestock Systems), hazard events (DroughtEvents, Floods), vulnerable demographics (Populations, LivelihoodZones), and institutional responses (Interventions, PSNP Cash/Food Transfers, EarlyWarningAlerts). Graph traversal queries enable complex contextual reasoning that enriches the decision-support interface.")

    # 2.7 Indigenous Ecological Knowledge
    add_heading_2(doc, "2.7 Indigenous Ecological Knowledge (IEK) and Agronomic Translation")
    add_body(doc, "Centuries before modern satellite remote sensing, smallholder farming and pastoralist communities in Eastern Ethiopia developed sophisticated, empirical Indigenous Ecological Knowledge (IEK) systems to forecast weather extremes and guide farming decisions. These indicators include atmospheric cues (shifts in prevailing dry easterly winds vs. moist southwesterlies; morning cloud formations over Harar and Gara Muleta mountains), faunal indicators (abnormal nesting and flight trajectories of the Abdim's Stork [Ciconia abdimii, locally known as 'Shimela']; intense honeybee swarming and aggressive foraging), and floral phenology (early flowering of Acacia bussei and Hagenia abyssinica).")
    add_body(doc, "Rather than ignoring or romanticizing these traditional cues, this research designs a formal, computational Bi-Directional IEK Consensus Engine. The engine cross-validates AI-derived probabilistic forecasts against community-submitted traditional bio-climatic observations, classifying the operational state into High Consensus, Moderate Consensus, or Divergent Consensus, thereby fostering reciprocal community ownership and elevating advisory adoption.")

    # 2.8 Conceptual Framework Blueprint
    add_heading_2(doc, "2.8 Conceptual Framework Blueprint")
    add_body(doc, "The overarching conceptual framework of this thesis is structured as a five-tiered operational hierarchy spanning from multi-source data ingestion to on-the-ground farmer action, as illustrated in the system architecture blueprint:")
    
    add_bullet(doc, "Automated, zero-credential connectors ingest global Pacific SST indices, ERA5 atmospheric reanalysis, CHIRPS v3 satellite rainfall, NASA SMAP soil moisture, NASA MODIS NDVI/LST, SoilGrids edaphic layers, WorldPop population grids, CSA agricultural statistics, and FEWS NET IPC classifications. Every ingestion run computes SHA-256 cryptographic manifests for immutable data provenance.", bold_prefix="Tier 1: Multi-Source Data Ingestion & Manifest Security: ")
    add_bullet(doc, "Data are aligned to authoritative Ethiopian administrative WGS84 woreda boundaries and an 8×8 topographic grid. Physical indices (Sc-PDSI, SPI-1/3/6, VCI, anomalies) are computed, and temporal multi-month sliding windows are assembled without look-ahead data leakage.", bold_prefix="Tier 2: Feature Engineering & Spatial Discretization: ")
    add_bullet(doc, "The Super-Hybrid Deep Learning Engine processes spatial patches (CNN), temporal sequences (LSTM), periodic harmonics (FFT), and structured tables (MLP), combining them via a Gated Cross-Attention Fusion Engine. Prediction heads emit Drought Hazard Probability, Crop Risk, and FEWS NET IPC Food Insecurity Phase.", bold_prefix="Tier 3: Super-Hybrid Multimodal Deep Learning Engine: ")
    add_bullet(doc, "Predictions are evaluated under the 'Honest Lead Horizon' principle (BSS gating). Crop hazard is mapped through biological sensitivity multipliers (s) onto a 4-tier prescriptive action matrix. The IEK consensus engine reconciles traditional bio-cues. XAI modules (SHAP, Grad-CAM) and the Neo4j Knowledge Graph provide human-interpretable contextual evidence.", bold_prefix="Tier 4: Agro-Ecological Decision Support, IEK Consensus & XAI: ")
    add_bullet(doc, "Delivers actionable intelligence through a cloud-native FastAPI backend (PostgreSQL 17 + Star Schema), a Next.js 16 Web Command Centre with role-scoped multi-tenancy (Ministers, Woreda Development Agents, Farmers), and an offline-first Expo React Native mobile application supporting English, Amharic, and Afaan Oromoo.", bold_prefix="Tier 5: Multi-Platform Enterprise Delivery Ecosystem: ")

    doc.add_page_break()

print("Chapter 2 module defined.")
